"""Genera web/data/evolucion.js (tabla de evolución que lee la app).

Entradas:
  datos/evolucion_wiktionary.csv  propuestas de étimo del Wikcionario en español (estado: pendiente)
  datos/verificacion_cid.xlsx     (opcional) tu revisión manual: Estado, Fuente verificada y étimo corregido

Reglas:
  - Solo se incluyen los lemas que figuran en web/data/cid.json.
  - Un lema queda «verificado» únicamente si en la hoja tiene Estado = verificado Y una Fuente verificada.
    Sin fuente se trata como pendiente (la IA no inventa etimologías; si una fuente no lo dice, la celda queda vacía).
  - Si corregís el étimo en la hoja, deja de valer la fuente del Wikcionario: se usa tu Fuente verificada.

Uso:  python scripts/build_evolucion.py
Datos del Wikcionario: CC BY-SA 4.0 (es.wiktionary.org, extraídos con wiktextract / kaikki.org).
"""
import csv, json, re, sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
CSV_WIKT = RAIZ / "datos" / "evolucion_wiktionary.csv"
XLSX = RAIZ / "datos" / "verificacion_cid.xlsx"
CID = RAIZ / "web" / "data" / "cid.json"
SALIDA = RAIZ / "web" / "data" / "evolucion.js"

CABECERA = """/* Tabla de evolución. Una entrada por lema (nombre del lema como en el DEMel, p. ej. "noche").
   Generado por scripts/build_evolucion.py: no editar a mano (se pisa); editar datos/verificacion_cid.xlsx.
   window.HYS_EVO = {
     "noche": {
       "cadena": [ {"forma":"noctem", "etapa":"latín"} ],
       "cambios": [], "marca": "patrimonial | semicultismo | cultismo | analogía",
       "fuentes": ["obra, página o entrada citada"],
       "estado": "verificado | pendiente",   // pendiente = propuesta sin verificar en la bibliografía
       "nota": "aviso opcional"
     }
   };
   Cada dato verificado debe tener su fuente. Propuestas del Wikcionario en español, CC BY-SA 4.0. */
"""

def cadena_de(etimo):
    partes = [p.strip() for p in etimo.split(" | ") if p.strip()]
    items = []
    for p in partes:
        m = re.match(r"^([a-záéíóúñ ]+): (.+)$", p)   # «árabe: قاض»; el latín va sin prefijo
        items.append((m.group(2), m.group(1)) if m else (p, "latín"))
    if len(items) > 1:
        return [{"forma": " o ".join(f for f, _ in items), "etapa": "varios étimos posibles"}]
    return [{"forma": items[0][0], "etapa": items[0][1]}] if items else []

def fuentes_wikt(texto):
    out = []
    for m in re.finditer(r"entrada «([^»]+)» \(([^)]*)\)([^:]*):", texto):
        s = "Wikcionario en español (es.wiktionary.org, CC BY-SA 4.0), entrada «%s» (%s)" % (m.group(1), m.group(2))
        if s not in out:
            out.append(s)
    return out

def main():
    cid = json.load(open(CID, encoding="utf-8"))
    lemas_cid = {v["n"] for v in cid["lemmas"].values()}
    base = {}
    for r in csv.DictReader(open(CSV_WIKT, encoding="utf-8-sig"), delimiter=";"):
        if r["etimo"].strip():
            base[r["lema"]] = r
    hoja = {}
    if XLSX.exists():
        try:
            from openpyxl import load_workbook
            for fila in load_workbook(XLSX, read_only=True, data_only=True)["Cid"].iter_rows(min_row=2, values_only=True):
                if fila[0]:
                    hoja[str(fila[0]).strip()] = [("" if c is None else str(c).strip()) for c in fila[:7]]
        except ImportError:
            print("AVISO: falta openpyxl (pip install openpyxl); se ignora la hoja de verificación.", file=sys.stderr)
    evo, avisos = {}, []
    for lema in sorted(lemas_cid | set(hoja)):
        if lema not in lemas_cid or re.fullmatch(r"[IVXLCDM]+", lema):   # sin números romanos (D, V…)
            continue
        b = base.get(lema)
        h = hoja.get(lema)
        etimo = b["etimo"].strip() if b else ""
        fuentes = fuentes_wikt(b["fuentes"]) if b else []
        nota = ("Hay más de un étimo posible (homónimos); falta elegir el que corresponde al sentido en el texto."
                if " | " in etimo else "")
        verificado = False
        if h:
            h_etimo, h_estado, h_fuente = h[1], h[4].lower(), h[5]
            if h_etimo and h_etimo != etimo:           # corrección o carga manual
                etimo, fuentes, nota = h_etimo, [], ""
            if h_fuente:
                fuentes = fuentes + [h_fuente] if h_etimo == (b["etimo"].strip() if b else "") else [h_fuente]
            if h_estado == "verificado":
                if h_fuente and etimo:
                    verificado = True
                else:
                    avisos.append("«%s» está marcado verificado pero sin Fuente verificada (o sin étimo): queda pendiente." % lema)
            if h[6]:
                nota = (nota + " " + h[6]).strip()
        if not etimo:
            continue
        if not fuentes:
            fuentes = ["Sin fuente indicada"]
        e = {"cadena": cadena_de(etimo), "cambios": [], "fuentes": fuentes,
             "estado": "verificado" if verificado else "pendiente"}
        if nota:
            e["nota"] = nota
        evo[lema] = e
    cuerpo = ",\n".join(" %s: %s" % (json.dumps(k, ensure_ascii=False), json.dumps(v, ensure_ascii=False)) for k, v in sorted(evo.items()))
    SALIDA.write_text(CABECERA + "window.HYS_EVO = {\n" + cuerpo + "\n};\n", encoding="utf-8", newline="\n")
    ver = sum(1 for v in evo.values() if v["estado"] == "verificado")
    print("evolucion.js: %d lemas con étimo (%d verificados, %d pendientes) de %d lemas del Cid" % (len(evo), ver, len(evo) - ver, len(lemas_cid)))
    for a in avisos:
        print("AVISO:", a, file=sys.stderr)

if __name__ == "__main__":
    main()
